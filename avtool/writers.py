"""Writing the output files.

Names keep the full source name: interview.mov -> interview.mov.srt, so
interview.mov and interview.mp4 in one folder never overwrite each other.

Every file is first written as "<name>.partial" and renamed only when
complete, so a crash never leaves a half-written transcript behind. The .json
is written LAST: it is the "this file is finished" marker that resume checks.
"""

from __future__ import annotations

import json
import os
from datetime import datetime
from pathlib import Path
from typing import Callable, Optional

from . import __version__
from .backends import Segment
from .text import CaptionStyle, Cue, Glossary, build_cues, paragraphs
from .timefmt import hms, human_duration, srt_time, vtt_time

PARTIAL = ".partial"
NO_SPEECH = "(No speech detected.)"

# format name -> file suffix added after "name.ext"
FORMATS = {
    "txt": ".txt",
    "timestamped": ".timestamped.txt",
    "srt": ".srt",
    "vtt": ".vtt",
    "docx": ".docx",
    "json": ".json",
}
DEFAULT_FORMATS = list(FORMATS)


def output_path(out_dir: Path, source: Path, fmt: str) -> Path:
    return out_dir / (source.name + FORMATS[fmt])


def done_marker(out_dir: Path, source: Path) -> Path:
    return output_path(out_dir, source, "json")


def remove_partials(out_dir: Path, source: Path) -> None:
    for fmt in FORMATS:
        final = output_path(out_dir, source, fmt)
        partial = final.with_name(final.name + PARTIAL)
        if partial.exists():
            partial.unlink()


def _atomic(path: Path, write: Callable[[Path], None]) -> None:
    tmp = path.with_name(path.name + PARTIAL)
    try:
        write(tmp)
        os.replace(tmp, path)  # the rename is all-or-nothing
    finally:
        if tmp.exists():
            tmp.unlink()


def _write_text(path: Path, text: str) -> None:
    def write(tmp: Path) -> None:
        with open(tmp, "w", encoding="utf-8", newline="\n") as f:
            f.write(text)
            f.flush()
            os.fsync(f.fileno())
    _atomic(path, write)


# ---------------------------------------------------------------- renderers

def render_txt(paras: list[tuple[float, str]]) -> str:
    if not paras:
        return NO_SPEECH + "\n"
    return "\n\n".join(p for _, p in paras) + "\n"


def render_timestamped(segments: list[Segment], glossary: Optional[Glossary]) -> str:
    if not segments:
        return NO_SPEECH + "\n"
    lines = []
    for seg in segments:
        text = glossary.apply(seg.text.strip()) if glossary else seg.text.strip()
        lines.append(f"[{hms(seg.start)}] {text}")
    return "\n".join(lines) + "\n"


def render_srt(cues: list[Cue]) -> str:
    blocks = [
        f"{i}\n{srt_time(c.start)} --> {srt_time(c.end)}\n" + "\n".join(c.lines)
        for i, c in enumerate(cues, 1)
    ]
    return "\n\n".join(blocks) + ("\n" if blocks else "")


def render_vtt(cues: list[Cue]) -> str:
    blocks = ["WEBVTT"]
    for c in cues:
        blocks.append(f"{vtt_time(c.start)} --> {vtt_time(c.end)}\n" + "\n".join(c.lines))
    return "\n\n".join(blocks) + "\n"


def write_docx(path: Path, title: str, meta: str, paras: list[tuple[float, str]]) -> None:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    doc.add_heading(title, level=1)
    info = doc.add_paragraph()
    run = info.add_run(meta)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    if not paras:
        doc.add_paragraph(NO_SPEECH)
    for start, text in paras:
        p = doc.add_paragraph()
        stamp = p.add_run(f"[{hms(start)}]  ")
        stamp.bold = True
        stamp.font.size = Pt(9)
        stamp.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        p.add_run(text)

    _atomic(path, lambda tmp: doc.save(str(tmp)))


# ---------------------------------------------------------------- all outputs

def write_all(
    out_dir: Path,
    source: Path,
    segments: list[Segment],
    *,
    formats: list[str],
    style: CaptionStyle,
    style_name: str,
    glossary: Optional[Glossary],
    language: Optional[str],
    model_label: str,
    duration: float,
    audio_track: int,
    settings: dict,
) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    paras = paragraphs(segments, glossary)
    cues = build_cues(segments, style, glossary)
    written: list[Path] = []

    if "txt" in formats:
        p = output_path(out_dir, source, "txt")
        _write_text(p, render_txt(paras))
        written.append(p)
    if "timestamped" in formats:
        p = output_path(out_dir, source, "timestamped")
        _write_text(p, render_timestamped(segments, glossary))
        written.append(p)
    if "srt" in formats:
        p = output_path(out_dir, source, "srt")
        _write_text(p, render_srt(cues))
        written.append(p)
    if "vtt" in formats:
        p = output_path(out_dir, source, "vtt")
        _write_text(p, render_vtt(cues))
        written.append(p)
    if "docx" in formats:
        p = output_path(out_dir, source, "docx")
        meta = (
            f"{human_duration(duration)} · transcribed {datetime.now():%b %d, %Y} · "
            f"{model_label} · language: {language or 'unknown'}"
        )
        try:
            write_docx(p, source.name, meta, paras)
            written.append(p)
        except Exception as exc:  # a Word problem shouldn't cost the other outputs
            print(f"  ⚠ Word file skipped ({exc}); the other files are fine.")

    # Last: the JSON. Its existence means "this file is completely done".
    data = {
        "tool": f"AVTool CMOS {__version__}",
        "created": datetime.now().isoformat(timespec="seconds"),
        "source": {
            "name": source.name,
            "path": str(source),
            "size_bytes": source.stat().st_size if source.exists() else None,
            "duration_seconds": round(duration, 3),
            "audio_track": audio_track,
        },
        "model": model_label,
        "language": language,
        "captions_preset": style_name,
        "settings": settings,
        "speech_found": bool(segments),
        "segments": [
            {
                "start": round(s.start, 3),
                "end": round(s.end, 3),
                "text": glossary.apply(s.text.strip()) if glossary else s.text.strip(),
                "avg_logprob": round(s.avg_logprob, 4),
                "no_speech_prob": round(s.no_speech_prob, 4),
                "words": [
                    {"start": round(w.start, 3), "end": round(w.end, 3),
                     "word": w.text, "probability": round(w.probability, 4)}
                    for w in s.words
                ],
            }
            for s in segments
        ],
        "captions": [
            {"start": round(c.start, 3), "end": round(c.end, 3), "lines": c.lines}
            for c in cues
        ],
    }
    p = done_marker(out_dir, source)
    _write_text(p, json.dumps(data, ensure_ascii=False, indent=2) + "\n")
    written.append(p)
    return written
